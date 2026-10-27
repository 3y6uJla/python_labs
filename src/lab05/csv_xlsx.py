import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

MIN_WIDTH = 8


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """Функция конвертирует CSV в XLSX

    Первая строка CSV - заголовок (выделяется жирным). Лист называется "Sheet1".
    Ширина каждой колонки подбирается по самому длинному значению, но не меньше 8 символов.

    Аргументы:
        csv_path: путь к входному CSV-файлу
        xlsx_path: путь к выходному XLSX-файлу

    Вывод:
        Ничего не возвращает, создаёт (или перезаписывает) XLSX-файл

    Ошибки:
        FileNotFoundError: если CSV-файла нет
        ValueError: неверный тип файла, пустой CSV, нет заголовка
            или строки разной длины
    """

    csv_file = Path(csv_path)
    xlsx_file = Path(xlsx_path)

    if csv_file.suffix.lower() != '.csv':
        raise ValueError(f'Неверный тип файла: {csv_file.name}, ожидается .csv')
    if xlsx_file.suffix.lower() != '.xlsx':
        raise ValueError(f'Неверный тип файла: {xlsx_file.name}, ожидается .xlsx')
    if not csv_file.is_file():
        raise FileNotFoundError(f'Файл не найден: {csv_file}')

    with open(csv_file, 'r', encoding='utf-8', newline='') as f:
        rows = list(csv.reader(f))

    if len(rows) == 0:
        raise ValueError('Пустой CSV')

    header = rows[0]
    if len(header) == 0 or '' in [name.strip() for name in header]:
        raise ValueError('В CSV нет заголовка или в нём есть пустая колонка')

    for i, row in enumerate(rows, start=1):
        if len(row) != len(header):
            raise ValueError(f'Неверная строка {i}: значений - {len(row)}, '
                             f'колонок в заголовке - {len(header)}')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'

    for row in rows:
        ws.append(row)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    for col in range(len(header)):
        max_len = max(len(row[col]) for row in rows)
        width = max(max_len, MIN_WIDTH) + 2
        ws.column_dimensions[get_column_letter(col + 1)].width = width

    xlsx_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_file)


def main() -> None:
    from openpyxl import load_workbook

    for src, dst in [('data/samples/people.csv', 'data/out/people.xlsx'),
                     ('data/samples/cities.csv', 'data/out/cities.xlsx')]:
        csv_to_xlsx(src, dst)
        ws = load_workbook(dst).active
        print(f'csv_to_xlsx("{src}", "{dst}")')
        print(f'  лист: {ws.title}, строк: {ws.max_row}, колонок: {ws.max_column}')
        widths = []
        for col in range(1, ws.max_column + 1):
            letter = get_column_letter(col)
            widths.append(f'{letter}={ws.column_dimensions[letter].width}')
        print(f'  ширина колонок: {", ".join(widths)}')

    print('\nПроверка ошибок:')
    for path in ['data/samples/empty.csv', 'data/samples/people.json',
                 'data/samples/bad_rows.csv', 'data/samples/nofile.csv']:
        try:
            csv_to_xlsx(path, 'data/out/bad.xlsx')
        except (ValueError, FileNotFoundError) as e:
            print(f'csv_to_xlsx("{path}", ...) -> {type(e).__name__}: {e}')


if __name__ == '__main__':
    main()
